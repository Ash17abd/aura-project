from __future__ import annotations

import io
import json
import os
from pathlib import Path
from typing import Any

import requests
import streamlit as st


APP_TITLE = "AURA online"
OPENROUTER_API_URL = "https://openrouter.ai/api/v1/chat/completions"
GEMINI_API_BASE = "https://generativelanguage.googleapis.com/v1beta/models"

DEFAULT_MODEL = "gemini-3.5-flash-lite"
MODELS = [
    "gemini-3.5-flash-lite",
    "gemini-3.5-flash",
    "google/gemma-3-27b-it:free",
    "meta-llama/llama-3.3-70b-instruct:free",
    "qwen/qwen3-coder:free",
]
MAX_CONTEXT_CHARS = 12000


def get_secret(name: str) -> str:
    """Retrieve secrets from Streamlit secrets, environment variables, or config."""
    # 1. Streamlit secrets
    try:
        val = st.secrets.get(name, "")
        if val:
            return str(val).strip()
    except Exception:
        pass

    # 2. Environment variables
    val = os.getenv(name, "")
    if val:
        return str(val).strip()

    # 3. Fallback to local config file if present
    try:
        config_path = Path(__file__).resolve().parent / "config" / "api_keys.json"
        if config_path.exists():
            with open(config_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            key_map = {
                "GEMINI_API_KEY": "gemini_api_key",
                "OPENROUTER_API_KEY": "openrouter_api_key",
            }
            mapped = key_map.get(name, name.lower())
            return str(data.get(mapped, "")).strip()
    except Exception:
        pass

    return ""


def extract_text(uploaded_file: Any) -> str:
    suffix = uploaded_file.name.lower().rsplit(".", 1)[-1]
    raw = uploaded_file.getvalue()

    if suffix == "pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
    if suffix == "docx":
        from docx import Document

        document = Document(io.BytesIO(raw))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == "pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(raw))
        lines: list[str] = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text)
        return "\n".join(lines)
    return raw.decode("utf-8", errors="replace")


def request_gemini_answer(messages: list[dict[str, str]], model: str) -> str:
    api_key = get_secret("GEMINI_API_KEY")
    if not api_key:
        # Fallback to OPENROUTER_API_KEY if it might be a Gemini key
        or_key = get_secret("OPENROUTER_API_KEY")
        if or_key and or_key.startswith("AQ."):
            api_key = or_key
    if not api_key:
        raise RuntimeError("Add GEMINI_API_KEY to environment variables or Streamlit secrets before chatting.")

    system_text = ""
    contents = []
    for m in messages:
        if m["role"] == "system":
            system_text = (system_text + "\n" + m["content"]).strip()
        else:
            role = "model" if m["role"] == "assistant" else "user"
            contents.append({"role": role, "parts": [{"text": m["content"]}]})

    payload: dict[str, Any] = {"contents": contents}
    if system_text:
        payload["system_instruction"] = {"parts": [{"text": system_text}]}

    url = f"{GEMINI_API_BASE}/{model}:generateContent?key={api_key}"
    response = requests.post(url, json=payload, timeout=90)
    if response.status_code in (503, 429) and model != "gemini-3.5-flash-lite":
        # Fallback to flash-lite on upstream demand spikes
        fallback_url = f"{GEMINI_API_BASE}/gemini-3.5-flash-lite:generateContent?key={api_key}"
        response = requests.post(fallback_url, json=payload, timeout=90)

    response.raise_for_status()
    data = response.json()
    candidates = data.get("candidates", [])
    if not candidates:
        raise RuntimeError("The model returned an empty response.")
    parts = candidates[0].get("content", {}).get("parts", [])
    answer = "".join(part.get("text", "") for part in parts if "text" in part)
    if not answer:
        raise RuntimeError("No text content returned from the model.")
    return answer.strip()


def request_openrouter_answer(messages: list[dict[str, str]], model: str) -> str:
    api_key = get_secret("OPENROUTER_API_KEY")
    if not api_key:
        raise RuntimeError("Add OPENROUTER_API_KEY to environment variables or Streamlit secrets before chatting.")

    response = requests.post(
        OPENROUTER_API_URL,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "HTTP-Referer": "https://streamlit.io/",
            "X-Title": APP_TITLE,
        },
        json={"model": model, "messages": messages, "temperature": 0.7, "max_tokens": 2048},
        timeout=90,
    )
    response.raise_for_status()
    data = response.json()
    answer = data.get("choices", [{}])[0].get("message", {}).get("content", "")
    if not answer:
        raise RuntimeError("The model returned an empty response.")
    return answer.strip()


def request_answer(messages: list[dict[str, str]], model: str) -> str:
    if model.startswith("gemini-"):
        return request_gemini_answer(messages, model)
    return request_openrouter_answer(messages, model)


st.set_page_config(page_title=APP_TITLE, page_icon=":material/auto_awesome:", layout="centered")

if "messages" not in st.session_state:
    st.session_state.messages = []
if "document_context" not in st.session_state:
    st.session_state.document_context = ""

st.title(APP_TITLE)
st.caption("A browser-friendly AURA workspace for conversation and study material.")

with st.sidebar:
    st.subheader("Workspace")
    model = st.selectbox(
        "Model",
        MODELS,
        index=0,
    )
    uploaded_file = st.file_uploader(
        "Add study material",
        type=["pdf", "docx", "pptx", "txt", "md", "csv", "json", "html", "py", "js", "ts"],
    )
    if uploaded_file is not None and st.button("Load document", icon=":material/upload_file:", use_container_width=True):
        try:
            text = extract_text(uploaded_file).strip()
            st.session_state.document_context = text[:MAX_CONTEXT_CHARS]
            st.success(f"Loaded {uploaded_file.name}")
        except Exception as exc:
            st.error(f"Could not read that file: {exc}")

    if st.session_state.document_context:
        st.caption(f"Document context: {len(st.session_state.document_context):,} characters")
    if st.button("Clear conversation", icon=":material/delete_sweep:", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

if not st.session_state.messages:
    st.info("Ask a question, summarize a document, or explore a topic.")

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

prompt = st.chat_input("Ask AURA anything")
if prompt:
    st.session_state.messages.append({"role": "user", "content": prompt})
    with st.chat_message("user"):
        st.markdown(prompt)

    system_prompt = (
        "You are AURA, a concise and capable learning assistant. "
        "Be accurate, practical, and transparent about uncertainty."
    )
    if st.session_state.document_context:
        system_prompt += (
            " Use the following uploaded document as context when relevant:\n\n"
            + st.session_state.document_context
        )
    api_messages = [{"role": "system", "content": system_prompt}, *st.session_state.messages]

    with st.chat_message("assistant"):
        with st.status("Thinking", type="compact"):
            try:
                answer = request_answer(api_messages, model)
            except requests.HTTPError as exc:
                answer = f"The model request failed: {exc.response.status_code}. Check your API key and model access."
            except Exception as exc:
                answer = str(exc)
        st.markdown(answer)
    st.session_state.messages.append({"role": "assistant", "content": answer})